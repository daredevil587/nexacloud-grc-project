# NexaCloud Ltd - GRC Risk Register Excel Builder
# Author: Uttam Yadav
# Creates a formatted Excel file from risk data

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
ws = wb.active
ws.title = "Risk Register"

headers = ["ID", "Risk Name", "Likelihood", "Impact", "Score", "Rating", "NIST Control", "Owner", "Status", "Mitigation"]
ws.append(headers)

for col in range(1, 11):
    cell = ws.cell(row=1, column=col)
    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="E25B2B")
    cell.alignment = Alignment(horizontal="center")

risks = [
    ["R-001", "Phishing Attack", 4, 4, 16, "HIGH", "IA-2, AT-2", "IT Security", "In Progress", "Enforce MFA on all accounts. Quarterly phishing training."],
    ["R-002", "Misconfigured S3 Bucket", 3, 5, 15, "HIGH", "CM-6, AU-2", "Cloud Engineering", "Planned", "Enable AWS Block Public Access."],
    ["R-003", "Ransomware Attack", 3, 5, 15, "HIGH", "CP-9, IR-4, SI-3", "IT Security", "In Progress", "Immutable backups tested monthly. EDR on all endpoints."],
    ["R-006", "Third-Party Vendor Breach", 3, 4, 12, "HIGH", "SR-3, SA-9", "GRC / Legal", "In Progress", "Annual vendor assessments."],
    ["R-007", "SQL Injection Vulnerability", 3, 4, 12, "HIGH", "SI-10, SI-2, SA-11", "Application Security", "In Progress", "Cloudflare WAF blocks OWASP Top 10."],
    ["R-009", "Shadow IT - Unapproved SaaS", 4, 3, 12, "HIGH", "CM-8, SA-9, AC-20", "IT Security", "Planned", "Cloudflare Gateway detects unapproved apps."],
    ["R-010", "Business Email Compromise", 3, 4, 12, "HIGH", "AT-2, SI-8", "Finance + IT Security", "In Progress", "DMARC/DKIM/SPF enforced."],
    ["R-004", "Insider Threat - Data Exfiltration", 2, 5, 10, "HIGH", "AC-6, AU-9, PS-4", "HR + IT Security", "Implemented", "Least privilege RBAC. 90-day access reviews."],
    ["R-005", "DDoS Platform Outage", 3, 3, 9, "MEDIUM", "SC-5, CP-8", "Platform Engineering", "Implemented", "Cloudflare Magic Transit enabled."],
    ["R-008", "GDPR Non-Compliance", 2, 4, 8, "MEDIUM", "PT-1, PT-2, PM-19", "DPO / Legal", "Implemented", "DSAR workflow implemented. DPO appointed."],
]

for row in risks:
    ws.append(row)

for row in range(2, 12):
    rating = ws.cell(row=row, column=6).value
    if rating == "HIGH":
        colour = "FAD4D4"
    elif rating == "MEDIUM":
        colour = "FFF3CD"
    else:
        colour = "D6EAD6"
    for col in range(1, 11):
        cell = ws.cell(row=row, column=col)
        cell.fill = PatternFill("solid", fgColor=colour)
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center")

ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 28
ws.column_dimensions["C"].width = 13
ws.column_dimensions["D"].width = 10
ws.column_dimensions["E"].width = 10
ws.column_dimensions["F"].width = 12
ws.column_dimensions["G"].width = 20
ws.column_dimensions["H"].width = 22
ws.column_dimensions["I"].width = 14
ws.column_dimensions["J"].width = 45

wb.save("risk_register.xlsx")
print("Excel file saved: risk_register.xlsx")